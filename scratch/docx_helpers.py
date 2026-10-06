import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''<w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>''')
        tblPr[0].append(borders)

def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=10, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=8, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=6, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_body_p(doc, text, bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False, space_after=4, font_size=10.5):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.18, align=align)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(font_size)
        r_pre.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.font.italic = italic
    return p

def add_bullet_p(doc, text, bold_prefix=None, space_after=3, font_size=10):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(font_size)
        r_pre.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    return p

def add_code_block(doc, code_text, space_after=6):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        <w:left w:val="single" w:sz="16" w:space="0" w:color="1E3A8A"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
    </w:tcBorders>''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=0, space_after=0, line_spacing=1.0)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    doc.add_paragraph()
    p_sp = doc.paragraphs[-1]
    format_paragraph(p_sp, space_before=0, space_after=space_after)

def add_callout_box(doc, title, text, border_color="1E3A8A", bg_color="EFF6FF", space_after=4):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    if bg_color:
        set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:bottom w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
        <w:right w:val="none"/>
    </w:tcBorders>''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=0, space_after=2, line_spacing=1.18)
    if title:
        r_t = p.add_run(title + "\n")
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(10.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(10)
    r_body.font.italic = True
    
    doc.add_paragraph()
    p_sp = doc.paragraphs[-1]
    format_paragraph(p_sp, space_before=0, space_after=space_after)

def add_template_rounded_box(doc, paragraphs_data, border_color="E07A22", bg_color=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    tcPr = cell._element.get_or_add_tcPr()
    if bg_color:
        set_cell_background(cell, bg_color)
    tcBorders = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="16" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="single" w:sz="16" w:space="0" w:color="{border_color}"/>
        <w:left w:val="single" w:sz="16" w:space="0" w:color="{border_color}"/>
        <w:right w:val="single" w:sz="16" w:space="0" w:color="{border_color}"/>
    </w:tcBorders>''')
    tcPr.append(tcBorders)
    
    for idx, (prefix, text, prefix_color) in enumerate(paragraphs_data):
        if idx == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        format_paragraph(p, space_before=0, space_after=4, line_spacing=1.18, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        if prefix:
            r_pre = p.add_run(prefix + " ")
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            if prefix_color:
                r_pre.font.color.rgb = prefix_color
        r_text = p.add_run(text)
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(10.5)
    
    doc.add_paragraph()
    p_sp = doc.paragraphs[-1]
    format_paragraph(p_sp, space_before=0, space_after=4)

def add_docx_watermark(section_or_header, image_path, width_in=5.5):
    if not os.path.exists(image_path):
        return
    if hasattr(section_or_header, 'header'):
        section = section_or_header
        header = section.header
    else:
        header = section_or_header
        section = None

    p = header.paragraphs[0]
    p.text = ""
    format_paragraph(p, space_before=0, space_after=0, line_spacing=1.0)
    r = p.add_run()
    pic = r.add_picture(image_path, width=Inches(width_in))
    inline = pic._inline
    
    graphic = inline.xpath('.//a:graphic')
    extent = inline.xpath('.//wp:extent')[0]
    cx = int(extent.get('cx'))
    cy = int(extent.get('cy'))
    
    if section is not None:
        page_w = int(section.page_width)
        page_h = int(section.page_height)
    else:
        page_w = int(Inches(8.27))
        page_h = int(Inches(11.69))
    
    pos_x = max(0, (page_w - cx) // 2)
    pos_y = max(0, (page_h - cy) // 2)
    
    anchor_xml = f'''<wp:anchor {nsdecls("wp")} {nsdecls("a")} {nsdecls("pic")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="0" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page">
            <wp:posOffset>{pos_x}</wp:posOffset>
        </wp:positionH>
        <wp:positionV relativeFrom="page">
            <wp:posOffset>{pos_y}</wp:posOffset>
        </wp:positionV>
        <wp:extent cx="{cx}" cy="{cy}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/>
        <wp:docPr id="666" name="Watermark"/>
        <wp:cNvGraphicFramePr/>
    </wp:anchor>'''
    anchor = parse_xml(anchor_xml)
    anchor.append(graphic[0])
    inline.getparent().replace(inline, anchor)

def add_centered_page_number(footer):
    p = footer.paragraphs[0]
    format_paragraph(p, space_before=4, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5)
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    run._r.append(fld)

print("Helper definitions successfully loaded.")

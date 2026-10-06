import docx
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
sec1 = doc.sections[0]
s1Pr = sec1._sectPr
pg1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
s1Pr.append(pg1)

p1 = doc.add_paragraph('Section 1 - Page 1')
doc.add_page_break()
p2 = doc.add_paragraph('Section 1 - Page 2')

sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
sec2.header.is_linked_to_previous = False
sec2.footer.is_linked_to_previous = False
s2Pr = sec2._sectPr
pg2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>')
s2Pr.append(pg2)

p3 = doc.add_paragraph('Section 2 - Page 1')

doc.save('scratch/test_numbering.docx')
print('Successfully saved test_numbering.docx')

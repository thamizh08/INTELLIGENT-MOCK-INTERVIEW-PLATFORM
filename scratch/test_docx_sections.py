import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
sec1 = doc.sections[0]
s1Pr = sec1._sectPr
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman"/>')
s1Pr.append(pgNumType)

p = doc.add_paragraph("This is frontmatter page 1")
doc.add_page_break()
p2 = doc.add_paragraph("This is frontmatter page 2")

sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
sec2.header.is_linked_to_previous = False
sec2.footer.is_linked_to_previous = False

s2Pr = sec2._sectPr
pgNumType2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>')
s2Pr.append(pgNumType2)

p3 = doc.add_paragraph("This is Chapter 1 page 1")

doc.save("scratch/test_sections_out.docx")
print("Saved test_sections_out.docx successfully!")

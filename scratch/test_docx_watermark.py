import docx
from docx import Document
from docx.shared import Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_docx_watermark(header, image_path, width_in=5.5):
    p = header.paragraphs[0]
    r = p.add_run()
    pic = r.add_picture(image_path, width=Inches(width_in))
    inline = pic._inline
    
    # Extract graphic from inline
    graphic = inline.xpath('.//a:graphic')
    extent = inline.xpath('.//wp:extent')[0]
    cx = extent.get('cx')
    cy = extent.get('cy')
    
    anchor_xml = f'''<wp:anchor {nsdecls("wp")} {nsdecls("a")} {nsdecls("pic")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page">
            <wp:align>center</wp:align>
        </wp:positionH>
        <wp:positionV relativeFrom="page">
            <wp:align>center</wp:align>
        </wp:positionV>
        <wp:extent cx="{cx}" cy="{cy}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/>
        <wp:docPr id="666" name="Watermark"/>
        <wp:cNvGraphicFramePr/>
    </wp:anchor>'''
    anchor = parse_xml(anchor_xml)
    anchor.append(graphic[0])
    inline.getparent().replace(inline, anchor)

doc = Document()
sec = doc.sections[0]
add_docx_watermark(sec.header, 'docs/extracted_assets/img_xref_46_387x113.jpeg')

# Add some body text to see if it renders over watermark
p = doc.add_paragraph("This is body text on top of watermark.")
doc.save('scratch/test_wm_behind.docx')
print("Successfully generated test_wm_behind.docx with behindDoc anchor!")
